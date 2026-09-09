#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Consulta automatica do EcoClass (Econet) para os NCM/NBS de uma planilha de
simulacao IBS/CBS e gera um arquivo .xlsx com o retorno completo da ferramenta.

Fluxo:
  1. Le os codigos NCM/NBS distintos das abas da planilha de origem.
  2. Autentica no portal Econet e obtem o token (JWT) da API do EcoClass.
  3. Para cada codigo chama POST /api/ecoclass/pesquisa e, quando ha detalhe,
     POST /api/ecoclass/obter-dados.
  4. Grava "Econet_EcoClass_Resultado.xlsx" (1+ linha por codigo) e um
     "Econet_EcoClass_bruto.jsonl" para auditoria.

Uso:
  python3 econet_ecoclass.py

Requisitos: requests, openpyxl  (pip install requests openpyxl)
"""

import json
import os
import re
import sys
import time
import urllib.parse
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------- #
# CONFIGURACAO
# --------------------------------------------------------------------------- #
BASE_DIR = Path(__file__).resolve().parent


def _carrega_dotenv(caminho=BASE_DIR / ".env"):
    """Carregador minimo de .env (sem dependencia externa)."""
    if not caminho.exists():
        return
    for linha in caminho.read_text(encoding="utf-8").splitlines():
        linha = linha.strip()
        if not linha or linha.startswith("#") or "=" not in linha:
            continue
        chave, _, valor = linha.partition("=")
        os.environ.setdefault(chave.strip(), valor.strip().strip('"').strip("'"))


_carrega_dotenv()

# Caminho da planilha de origem: use a variavel de ambiente ECOCLASS_WORKBOOK
# ou deixe o .xlsx na mesma pasta (o primeiro "Simulacao_*.xlsx" encontrado).
_wb_env = os.environ.get("ECOCLASS_WORKBOOK")
if _wb_env:
    WORKBOOK = Path(_wb_env)
else:
    _achados = sorted(BASE_DIR.glob("Simulacao_*.xlsx"))
    WORKBOOK = _achados[0] if _achados else BASE_DIR / "Simulacao.xlsx"

OUT_XLSX = BASE_DIR / "Econet_EcoClass_Resultado.xlsx"
OUT_RAW = BASE_DIR / "Econet_EcoClass_bruto.jsonl"

# Credenciais do Econet -> variaveis de ambiente / arquivo .env (nao versionado).
# Veja .env.example.
ECONET_LOGIN = os.environ.get("ECONET_LOGIN", "")
ECONET_SENHA = os.environ.get("ECONET_SENHA", "")

# Colunas de cada aba (1-indexado, como no Excel):
#   NFCe / NFe_Saidas -> coluna M = NCM
#   NFe_Entradas      -> coluna N = NCM
#   NFSe_Entradas     -> coluna P = NBS
SHEETS = [
    {"aba": "NFCe",          "col": "M", "header_row": 1, "tipo": "NCM"},
    {"aba": "NFe_Saidas",    "col": "M", "header_row": 1, "tipo": "NCM"},
    {"aba": "NFe_Entradas",  "col": "N", "header_row": 1, "tipo": "NCM"},
    {"aba": "NFSe_Entradas", "col": "P", "header_row": 4, "tipo": "NBS"},
]

PAUSA_ENTRE_CHAMADAS = 0.4  # segundos (educado com o servidor)
MAX_TENTATIVAS = 4

# --------------------------------------------------------------------------- #
# ENDPOINTS
# --------------------------------------------------------------------------- #
UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36")
APP_ORIGIN = "https://app.econeteditora.com.br"
LOGIN_ASP = "https://www.econeteditora.com.br/user/login.asp"
VERLOG_ASP = "https://www.econeteditora.com.br/user/ver_log.asp"
GO_LOGIN = "https://go-login.econeteditora.com.br/api/login"
CAPTCHA_GEN = "https://1u9nevrqx5.execute-api.sa-east-1.amazonaws.com/default/go-gerar-captcha"
CAPTCHA_XKEY = "4XQdAIUfrJ1Zpk2Jh0uRyaD5cNywcxQA9E5pvjJI"
ECOCLASS_API = "https://ecoclass-api.econeteditora.com.br/api/ecoclass"
PAG_DESTINO = "https://app.econeteditora.com.br/app/eco-class"

TIPO_API = {"NCM": "produto_mercadoria", "NBS": "servico"}


# --------------------------------------------------------------------------- #
# 1. EXTRACAO DOS CODIGOS
# --------------------------------------------------------------------------- #
def _limpa_codigo(valor, tipo):
    """Normaliza a celula -> string de digitos (ou None se nao for codigo)."""
    if valor is None:
        return None
    txt = str(valor).strip()
    if not txt:
        return None
    # Ex.: "115090000 - Servicos de processamento de dados" -> "115090000"
    m = re.match(r"\s*(\d[\d.]*\d|\d)", txt)
    if not m:
        return None
    digitos = re.sub(r"\D", "", m.group(1))
    if not digitos:
        return None
    # descarta lixo tipo "999999999" / "0000000"
    if len(set(digitos)) == 1:
        return None
    if len(digitos) < 4:
        return None
    return digitos


def extrair_codigos(caminho_wb):
    print(f"Lendo {caminho_wb.name} ...")
    wb = openpyxl.load_workbook(caminho_wb, read_only=True, data_only=True)
    registro = OrderedDict()  # (tipo, codigo) -> {"fontes": {aba: qtd}, "amostras": set}
    for cfg in SHEETS:
        aba = cfg["aba"]
        if aba not in wb.sheetnames:
            print(f"  [aviso] aba '{aba}' nao encontrada, ignorando.")
            continue
        ws = wb[aba]
        col_idx = openpyxl.utils.column_index_from_string(cfg["col"]) - 1
        # coluna de descricao (uma antes do NCM em NFCe/NFe; em NFSe usa "Descricao do Servico")
        desc_idx = col_idx - 1
        vistos = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i < cfg["header_row"]:
                continue
            if col_idx >= len(row):
                continue
            cod = _limpa_codigo(row[col_idx], cfg["tipo"])
            if not cod:
                continue
            chave = (cfg["tipo"], cod)
            reg = registro.setdefault(chave, {"fontes": OrderedDict(), "amostras": []})
            reg["fontes"][aba] = reg["fontes"].get(aba, 0) + 1
            if len(reg["amostras"]) < 3 and 0 <= desc_idx < len(row) and row[desc_idx]:
                d = str(row[desc_idx]).strip()
                if d and d not in reg["amostras"]:
                    reg["amostras"].append(d)
            vistos += 1
        print(f"  {aba}: {vistos} itens com {cfg['tipo']} valido")
    wb.close()

    codigos = []
    for (tipo, cod), reg in registro.items():
        codigos.append({
            "tipo": tipo,
            "codigo": cod,
            "fontes": reg["fontes"],
            "total_itens": sum(reg["fontes"].values()),
            "amostras": reg["amostras"],
        })
    codigos.sort(key=lambda c: (c["tipo"], c["codigo"]))
    print(f"Total de codigos distintos: {len(codigos)} "
          f"({sum(1 for c in codigos if c['tipo']=='NCM')} NCM, "
          f"{sum(1 for c in codigos if c['tipo']=='NBS')} NBS)\n")
    return codigos


# --------------------------------------------------------------------------- #
# 2. CLIENTE ECONET / ECOCLASS
# --------------------------------------------------------------------------- #
class EcoClassClient:
    def __init__(self, login, senha):
        self.login = login
        self.senha = senha
        self.s = requests.Session()
        self.s.headers.update({"User-Agent": UA, "Accept-Language": "pt-BR,pt;q=0.9"})
        self.jwt = None

    def autenticar(self):
        self.s.get(LOGIN_ASP, params={"Pag": PAG_DESTINO}, timeout=30)
        self.s.get(VERLOG_ASP, params={
            "Pag": PAG_DESTINO, "Origem": "Comex",
            "Log": self.login, "Sen": self.senha, "g-recaptcha-response": "",
        }, timeout=30)
        raw = self.s.cookies.get("bG9naW4", domain=".econeteditora.com.br")
        if not raw:
            raise RuntimeError("Login falhou: cookie de sessao nao retornado "
                               "(verifique login/senha).")
        cookie = urllib.parse.unquote(raw)
        r = self.s.post(GO_LOGIN, json={"cookie": cookie},
                        headers={"Origin": APP_ORIGIN, "Referer": APP_ORIGIN + "/"},
                        timeout=30)
        if r.status_code not in (200, 201):
            raise RuntimeError(f"Troca de token falhou: {r.status_code} {r.text[:200]}")
        self.jwt = r.text.strip().strip('"')
        if not self.jwt or self.jwt.count(".") != 2:
            raise RuntimeError(f"Token JWT invalido: {self.jwt[:80]}")
        print("Autenticado no EcoClass.\n")

    def _captcha(self, target_url):
        r = self.s.get(CAPTCHA_GEN, params={"path": target_url},
                       headers={"x-api-key": CAPTCHA_XKEY, "Referer": APP_ORIGIN + "/"},
                       timeout=30)
        r.raise_for_status()
        return r.json()["token"]

    def _post(self, endpoint, body):
        url = f"{ECOCLASS_API}/{endpoint}"
        ultimo_erro = None
        for tentativa in range(1, MAX_TENTATIVAS + 1):
            try:
                tok = self._captcha(url)
                r = self.s.post(url, json=body, headers={
                    "Authorization": self.jwt,
                    "g-recaptcha-response": tok,
                    "Origin": APP_ORIGIN, "Referer": APP_ORIGIN + "/",
                    "Content-Type": "application/json",
                }, timeout=45)
                if r.status_code == 401:  # token expirou
                    print("  token expirado, reautenticando ...")
                    self.autenticar()
                    continue
                r.raise_for_status()
                return r.json()
            except Exception as exc:  # noqa: BLE001
                ultimo_erro = exc
                time.sleep(1.5 * tentativa)
        raise RuntimeError(f"{endpoint} falhou apos {MAX_TENTATIVAS} tentativas: {ultimo_erro}")

    def pesquisa(self, tipo, codigo):
        return self._post("pesquisa", {"tipo": TIPO_API[tipo], "pesquisa": codigo})

    def obter_dados(self, tipo, item_id, id_cclass, id_ccredpres,
                    tributacao_integral, nbs_codigo=None):
        return self._post("obter-dados", {
            "id": item_id,
            "id_cclasstrib_cbs_ibs": id_cclass,
            "id_ccredpres": id_ccredpres,
            "tipo": TIPO_API[tipo],
            "tributacao_integral": bool(tributacao_integral),
            "nbs_codigo": nbs_codigo,
        })


# --------------------------------------------------------------------------- #
# 3. NORMALIZACAO DO RETORNO -> LINHAS DA PLANILHA
# --------------------------------------------------------------------------- #
def _join(seq, sep=" | "):
    return sep.join(x for x in seq if x)


def _bl(base_legal):
    if not base_legal:
        return "", ""
    txt = _join(b.get("base_legal", "") for b in base_legal)
    lnk = _join(b.get("link", "") for b in base_legal)
    return txt, lnk


def _sim_nao(v):
    if v in (1, True, "1"):
        return "Sim"
    if v in (0, False, "0", None):
        return "Nao"
    return str(v)


def linha_base(cod):
    return OrderedDict([
        ("Tipo", cod["tipo"]),
        ("Codigo consultado", cod["codigo"]),
        ("Abas de origem", _join(f"{a}({q})" for a, q in cod["fontes"].items())),
        ("Qtd. itens na planilha", cod["total_itens"]),
        ("Descricao (planilha - amostra)", _join(cod["amostras"], " /// ")),
        ("Status retorno", ""),
        ("NCM/NBS retornado", ""),
        ("Descricao (LC 214/2025)", ""),
        ("Tipo de beneficio", ""),
        ("CST", ""),
        ("CST - descricao", ""),
        ("cClassTrib", ""),
        ("cClassTrib - descricao", ""),
        ("Reducao de aliquota", ""),
        ("Reducao base de calculo", ""),
        ("Tributacao regular", ""),
        ("Credito presumido", ""),
        ("cCredPres", ""),
        ("cCredPres - descricao", ""),
        ("% CBS", ""),
        ("% IBS", ""),
        ("Tipo de aliquota", ""),
        ("Tratamento tributario", ""),
        ("Estorno de credito", ""),
        ("Inicio vigencia", ""),
        ("Fim vigencia", ""),
        ("Enquadramento / segmento", ""),
        ("Documentos fiscais", ""),
        ("Aliquota CBS vigente (tabela Econet)", ""),
        ("Aliquota IBS vigente (tabela Econet)", ""),
        ("Base legal", ""),
        ("Base legal - links", ""),
        ("JSON pesquisa", ""),
        ("JSON obter-dados", ""),
    ])


def processa_codigo(client, cod):
    """Retorna lista de OrderedDict (uma ou mais linhas)."""
    tipo = cod["tipo"]
    try:
        resp = client.pesquisa(tipo, cod["codigo"])
    except Exception as exc:  # noqa: BLE001
        l = linha_base(cod)
        l["Status retorno"] = f"ERRO: {exc}"
        return [l]

    data = (resp or {}).get("data", {}) or {}
    chave = "ncm" if tipo == "NCM" else "nbs"
    itens = data.get(chave) or []

    # CST puro (busca por 3 digitos) - nao esperado aqui, mas trata
    if not itens and isinstance(data.get("cst"), dict):
        l = linha_base(cod)
        cst = data["cst"]
        l["Status retorno"] = "Retorno de CST (codigo de 3 digitos)"
        l["CST"] = cst.get("codigo", "")
        l["CST - descricao"] = cst.get("descricao", "")
        bl, lk = _bl(cst.get("base_legal"))
        l["Base legal"], l["Base legal - links"] = bl, lk
        l["JSON pesquisa"] = json.dumps(resp, ensure_ascii=False)
        return [l]

    # Nao encontrado -> regra da ferramenta: CST 000 Tributacao Integral
    if not itens:
        l = linha_base(cod)
        l["Status retorno"] = "Nao localizado na LC 214/2025 -> CST 000 (Tributacao Integral)"
        l["CST"] = "000"
        l["CST - descricao"] = "Tributacao integral"
        l["cClassTrib"] = "000001"
        l["cClassTrib - descricao"] = "Situacoes tributadas integralmente pelo IBS e CBS."
        l["% CBS"] = "100.00"
        l["% IBS"] = "100.00"
        l["JSON pesquisa"] = json.dumps(resp, ensure_ascii=False)
        return [l]

    linhas = []
    for item in itens:
        l = linha_base(cod)
        l["NCM/NBS retornado"] = item.get("ncm") or item.get("nbs") or ""
        l["Descricao (LC 214/2025)"] = item.get("descricao", "")
        l["Tipo de beneficio"] = item.get("tipo_beneficio") or ""
        integral = bool(item.get("tributacao_integral"))
        item_id = item.get("id")
        id_cclass = item.get("id_cclasstrib_cbs_ibs")
        id_ccredpres = item.get("id_ccredpres")
        bl, lk = _bl(item.get("base_legal"))
        l["Base legal"], l["Base legal - links"] = bl, lk
        l["Inicio vigencia"] = item.get("inicio_vigencia") or ""
        l["Fim vigencia"] = item.get("fim_vigencia") or ""
        l["% CBS"] = item.get("valor_cbs") or ""
        l["% IBS"] = item.get("valor_ibs") or ""
        enq = item.get("enquadramento") or []
        l["Enquadramento / segmento"] = _join(
            f"{e.get('ncm_subitem','')}: {e.get('descricao','')}".strip(": ") for e in enq)

        cc = item.get("cclasstrib_cbs_ibs") or {}
        if cc:
            l["cClassTrib"] = cc.get("cclasstrib", "")
            l["cClassTrib - descricao"] = cc.get("descricao", "")
            l["Reducao de aliquota"] = _sim_nao(cc.get("reducao_aliquota"))
            l["Reducao base de calculo"] = _sim_nao(cc.get("reducao_base"))
            l["Tributacao regular"] = _sim_nao(cc.get("tributacao_regular", cc.get("tributacao_aliquota")))
            l["Credito presumido"] = _sim_nao(cc.get("credito_presumido"))
            l["Tipo de aliquota"] = cc.get("tipo_aliquota", "")
            l["Estorno de credito"] = _sim_nao(cc.get("estorno_credito"))

        # Detalhe (obter-dados) -------------------------------------------------
        det = None
        if item_id is not None and (id_cclass is not None or integral):
            try:
                nbs_cod = l["NCM/NBS retornado"].replace(".", "") if tipo == "NBS" else None
                det = client.obter_dados(tipo, item_id, id_cclass, id_ccredpres,
                                         integral, nbs_cod)
            except Exception as exc:  # noqa: BLE001
                l["JSON obter-dados"] = f"ERRO: {exc}"

        if det:
            l["JSON obter-dados"] = json.dumps(det, ensure_ascii=False)
            dd = (det or {}).get("data", {}) or {}
            cclist = dd.get("cClassTrib") or []
            if cclist:
                c0 = cclist[0]
                l["cClassTrib"] = c0.get("cclasstrib", l["cClassTrib"])
                l["cClassTrib - descricao"] = c0.get("descricao", l["cClassTrib - descricao"])
                l["Reducao de aliquota"] = _sim_nao(c0.get("reducao_aliquota"))
                l["Reducao base de calculo"] = _sim_nao(c0.get("reducao_base"))
                l["Tributacao regular"] = _sim_nao(c0.get("tributacao_regular"))
                l["Credito presumido"] = _sim_nao(c0.get("credito_presumido"))
                l["Tipo de aliquota"] = c0.get("tipo_aliquota", l["Tipo de aliquota"])
                l["Estorno de credito"] = _sim_nao(c0.get("estorno_credito"))
                cstd = c0.get("cst_cbs_ibs") or {}
                if cstd:
                    l["CST"] = cstd.get("codigo", l["CST"])
                    l["CST - descricao"] = cstd.get("descricao", l["CST - descricao"])
                tt = c0.get("tratamento_tributario") or {}
                if tt:
                    l["Tratamento tributario"] = _join([
                        tt.get("tratamento_tributario", ""), tt.get("tipo_tratamento", "")], " - ")
                if not l["Inicio vigencia"]:
                    l["Inicio vigencia"] = c0.get("inicio_vigencia") or ""
                if not l["Fim vigencia"]:
                    l["Fim vigencia"] = c0.get("fim_vigencia") or ""
                dfe = c0.get("documentos_fiscais_eletronicos") or []
                l["Documentos fiscais"] = _join(d.get("sigla", "") for d in dfe)
                if not l["Base legal"]:
                    l["Base legal"], l["Base legal - links"] = _bl(c0.get("base_legal"))
            cstlist = dd.get("cst") or []
            if cstlist and not l["CST"]:
                l["CST"] = cstlist[0].get("codigo", "")
                l["CST - descricao"] = cstlist[0].get("descricao", "")
            ali = dd.get("aliquota") or {}
            if isinstance(ali, dict):
                if isinstance(ali.get("cbs"), dict):
                    l["Aliquota CBS vigente (tabela Econet)"] = ali["cbs"].get("aliquota", "")
                if isinstance(ali.get("ibs"), dict):
                    l["Aliquota IBS vigente (tabela Econet)"] = ali["ibs"].get("aliquota", "")
            ccp = dd.get("cCredPres") or []
            if ccp:
                l["cCredPres"] = _join(str(x.get("ccredpres", "")) for x in ccp)
                l["cCredPres - descricao"] = _join(x.get("descricao", "") for x in ccp)
                l["Credito presumido"] = "Sim"

        # Status --------------------------------------------------------------
        if integral and not l["Tipo de beneficio"]:
            l["Status retorno"] = "Tributacao integral"
            if not l["CST"]:
                l["CST"], l["CST - descricao"] = "000", "Tributacao integral"
            if not l["cClassTrib"]:
                l["cClassTrib"] = "000001"
            if not l["% CBS"]:
                l["% CBS"] = "100.00"
            if not l["% IBS"]:
                l["% IBS"] = "100.00"
        else:
            l["Status retorno"] = f"Beneficio: {l['Tipo de beneficio'] or 'ver cClassTrib'}"

        l["JSON pesquisa"] = json.dumps({"data": {chave: [item]}}, ensure_ascii=False)
        linhas.append(l)
        time.sleep(PAUSA_ENTRE_CHAMADAS)
    return linhas


# --------------------------------------------------------------------------- #
# 4. GRAVACAO DO XLSX
# --------------------------------------------------------------------------- #
def grava_xlsx(linhas, caminho):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EcoClass"
    cols = list(linhas[0].keys())
    ws.append(cols)

    head_fill = PatternFill("solid", fgColor="1F4E78")
    head_font = Font(bold=True, color="FFFFFF")
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for l in linhas:
        ws.append([l.get(c, "") for c in cols])

    # cores por status
    fill_ok = PatternFill("solid", fgColor="E2EFDA")     # verde - beneficio
    fill_int = PatternFill("solid", fgColor="FFF2CC")    # amarelo - integral
    fill_err = PatternFill("solid", fgColor="FCE4D6")    # vermelho - erro/nao localizado
    st_idx = cols.index("Status retorno") + 1
    for r in range(2, ws.max_row + 1):
        st = str(ws.cell(row=r, column=st_idx).value or "")
        if st.startswith("ERRO") or "Nao localizado" in st:
            fill = fill_err
        elif "integral" in st.lower():
            fill = fill_int
        elif st.startswith("Beneficio"):
            fill = fill_ok
        else:
            fill = None
        if fill:
            for c in range(1, len(cols) + 1):
                ws.cell(row=r, column=c).fill = fill

    larguras = {
        "Descricao (planilha - amostra)": 42, "Descricao (LC 214/2025)": 60,
        "cClassTrib - descricao": 60, "CST - descricao": 26, "Base legal": 45,
        "Base legal - links": 45, "Tratamento tributario": 30,
        "Enquadramento / segmento": 40, "Abas de origem": 28,
        "JSON pesquisa": 40, "JSON obter-dados": 40, "cCredPres - descricao": 40,
    }
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(c, 16)

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"

    # aba de resumo
    rs = wb.create_sheet("Resumo")
    total = len(linhas)
    ben = sum(1 for l in linhas if str(l["Status retorno"]).startswith("Beneficio"))
    integ = sum(1 for l in linhas if "integral" in str(l["Status retorno"]).lower())
    nloc = sum(1 for l in linhas if "Nao localizado" in str(l["Status retorno"]))
    err = sum(1 for l in linhas if str(l["Status retorno"]).startswith("ERRO"))
    rs.append(["Consulta EcoClass - Econet", ""])
    rs.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
    rs.append(["Planilha de origem", WORKBOOK.name])
    rs.append(["Login Econet", ECONET_LOGIN])
    rs.append([])
    rs.append(["Linhas de resultado", total])
    rs.append(["  com beneficio", ben])
    rs.append(["  tributacao integral", integ])
    rs.append(["  nao localizado (CST 000)", nloc])
    rs.append(["  com erro", err])
    rs.append([])
    rs.append(["ATENCAO: o retorno reflete a classificacao geral da LC 214/2025. "
               "A empresa atua no segmento de restaurantes/alimentacao, que possui "
               "regra especifica - os CST/cClassTrib abaixo devem ser analisados "
               "criticamente antes de qualquer uso fiscal.", ""])
    rs.column_dimensions["A"].width = 40
    rs.column_dimensions["B"].width = 50

    wb.save(caminho)
    print(f"\nArquivo gravado: {caminho}")
    print(f"  {total} linhas | {ben} beneficio | {integ} integral | "
          f"{nloc} nao localizado | {err} erro")


# --------------------------------------------------------------------------- #
# MAIN
# --------------------------------------------------------------------------- #
def main():
    if not ECONET_LOGIN or not ECONET_SENHA:
        sys.exit("Defina ECONET_LOGIN e ECONET_SENHA (variavel de ambiente ou arquivo .env). "
                 "Veja .env.example.")
    if not WORKBOOK.exists():
        sys.exit(f"Planilha nao encontrada: {WORKBOOK}  "
                 "(defina ECOCLASS_WORKBOOK ou coloque o Simulacao_*.xlsx nesta pasta)")

    codigos = extrair_codigos(WORKBOOK)
    if not codigos:
        sys.exit("Nenhum codigo NCM/NBS encontrado nas abas configuradas.")

    client = EcoClassClient(ECONET_LOGIN, ECONET_SENHA)
    client.autenticar()

    todas_linhas = []
    with open(OUT_RAW, "w", encoding="utf-8") as fraw:
        for n, cod in enumerate(codigos, start=1):
            print(f"[{n:>3}/{len(codigos)}] {cod['tipo']} {cod['codigo']} "
                  f"({cod['total_itens']} itens) ...", end=" ", flush=True)
            linhas = processa_codigo(client, cod)
            for l in linhas:
                fraw.write(json.dumps({
                    "codigo": cod["codigo"], "tipo": cod["tipo"],
                    "status": l["Status retorno"],
                    "json_pesquisa": l["JSON pesquisa"],
                    "json_obter_dados": l["JSON obter-dados"],
                }, ensure_ascii=False) + "\n")
                fraw.flush()
            todas_linhas.extend(linhas)
            print(linhas[0]["Status retorno"][:60])
            time.sleep(PAUSA_ENTRE_CHAMADAS)

    grava_xlsx(todas_linhas, OUT_XLSX)
    print(f"JSON bruto: {OUT_RAW}")


if __name__ == "__main__":
    main()
